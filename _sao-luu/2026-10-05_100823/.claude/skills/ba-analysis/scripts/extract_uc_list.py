#!/usr/bin/env python3
"""
Lượt 1 — Trích xuất danh sách UC từ file Excel ra uc-inventory.md.

Nguyên tắc: chép nguyên văn mọi cột, không tóm tắt, không dịch; mỗi UC ghi vị trí nguồn.
Kèm kiểm kê nguồn (sheet/hàng/cột ẩn, comment, hình, công thức) và kiểm tra vệ sinh dữ liệu
(ID trùng, ID có khoảng trắng, ID nhảy số, ô trống ở cột chính, tên UC trùng actor).

Dùng:
  python extract_uc_list.py <file.xlsx> <out.md> [--annotations ann.json] [--summary summary.json]

ann.json (tùy chọn) — trường suy ra do bước phân tích thêm vào, hiển thị in nghiêng:
  {"UC-AUTH-001": {"Nhóm phạm vi": "A", "Phát hiện": "F-01"}}
"""
import argparse, json, re, sys, collections
from openpyxl import load_workbook

ID_RE = re.compile(r"^UC-[A-Z]+-\d+")
SECTION_RE = re.compile(r"^\d+\.\s")
KEY_COLS = ("name", "actor", "desc")


def norm(v):
    if v is None:
        return None
    s = str(v).replace("\r", "").strip()
    s = re.sub(r"\s*\n\s*", " / ", s)
    return s or None


def find_col(headers, *keys):
    for i, h in enumerate(headers):
        hl = (h or "").lower()
        if any(k in hl for k in keys):
            return i
    return None


def misaligned(headers, ucs):
    """Cờ lệch cột: giá trị của một cột trùng với giá trị 'danh mục' thường gặp ở cột khác."""
    cats = collections.defaultdict(collections.Counter)
    for uc in ucs:
        for h, v in uc["fields"]:
            if v and len(v) <= 60:
                cats[h][v] += 1
    common = {h: {v for v, n in c.items() if n >= 3} for h, c in cats.items()}
    for uc in ucs:
        for h, v in uc["fields"]:
            if not v or v in common.get(h, set()):
                continue
            for other, vals in common.items():
                if other != h and v in vals:
                    uc["flags"].append(f"Giá trị cột '{h}' giống giá trị thường gặp ở cột '{other}' (có thể lệch cột)")
                    break


def inventory_source(wb):
    rep = []
    for ws in wb.worksheets:
        hidden_rows = [r for r, d in ws.row_dimensions.items() if d.hidden]
        hidden_cols = [c for c, d in ws.column_dimensions.items() if d.hidden]
        comments = sum(1 for row in ws.iter_rows() for c in row if c.comment)
        formulas = sum(1 for row in ws.iter_rows() for c in row
                       if isinstance(c.value, str) and c.value.startswith("="))
        rep.append(dict(sheet=ws.title, state=ws.sheet_state, dims=ws.dimensions,
                        hidden_rows=hidden_rows, hidden_cols=hidden_cols,
                        merged=len(ws.merged_cells.ranges), comments=comments,
                        images=len(getattr(ws, "_images", [])), charts=len(getattr(ws, "_charts", [])),
                        formulas=formulas))
    return rep


def extract(path):
    wb = load_workbook(path)  # không data_only: cần thấy công thức
    sheets = []
    for ws in wb.worksheets:
        headers = [norm(ws.cell(1, c).value) for c in range(1, ws.max_column + 1)]
        cols = dict(name=find_col(headers, "use case name", "name"),
                    actor=find_col(headers, "actor"),
                    desc=find_col(headers, "description"))
        cur_section, ucs, sections, others = None, [], [], []
        for r in range(2, ws.max_row + 1):
            raw = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
            if all(v is None for v in raw):
                continue
            first = raw[0]
            first_s = str(first) if first is not None else ""
            if ID_RE.match(first_s.strip()):
                vals = [norm(v) for v in raw]
                uc = dict(id=first_s.strip(), row=r, sheet=ws.title, section=cur_section,
                          fields=list(zip(headers, vals)),
                          flags=[])
                if first_s != first_s.strip():
                    uc["flags"].append("ID có khoảng trắng thừa")
                for k in KEY_COLS:
                    i = cols[k]
                    if i is not None and not vals[i]:
                        uc["flags"].append(f"Cột '{headers[i]}' trống")
                if cols["name"] is not None and cols["actor"] is not None and vals[cols["name"]] \
                        and vals[cols["name"]] == vals[cols["actor"]]:
                    uc["flags"].append("Tên UC trùng nội dung cột Actor")
                ucs.append(uc)
            elif SECTION_RE.match(first_s.strip()):
                cur_section = first_s.strip()
                extra = [(h, norm(v)) for h, v in list(zip(headers, raw))[1:] if v is not None]
                sections.append((r, cur_section, extra))
            else:
                others.append((r, [(h, norm(v)) for h, v in zip(headers, raw) if v is not None]))
        misaligned(headers, ucs)
        sheets.append(dict(title=ws.title, headers=headers, ucs=ucs, sections=sections, others=others))
    return wb, sheets


def rare_actors(sheets):
    """Cờ actor hiếm: token vai trò chỉ xuất hiện 1 lần trong toàn file → soát chính tả/tên vai trò."""
    toks = collections.Counter()
    per_uc = []
    for sh in sheets:
        ai = next((i for i, h in enumerate(sh["headers"]) if h and "actor" in h.lower()), None)
        if ai is None:
            continue
        for uc in sh["ucs"]:
            v = uc["fields"][ai][1] or ""
            ts = [t.strip(" —-") for t in re.split(r"[,()+;]", v) if t.strip(" —-")]
            ts = [t for t in ts if not re.search(r"read-only|view only|assigned|including", t, re.I)]
            toks.update(ts)
            per_uc.append((uc, ts))
    for uc, ts in per_uc:
        for t in ts:
            if toks[t] == 1:
                uc["flags"].append(f"Actor hiếm gặp (1 lần trong file): '{t}' — soát chính tả/tên vai trò")


def hygiene(sheets):
    ids = collections.defaultdict(list)
    for sh in sheets:
        for uc in sh["ucs"]:
            if uc["id"]:
                ids[uc["id"]].append(f'{sh["title"]}!R{uc["row"]}')
    dup = {k: v for k, v in ids.items() if len(v) > 1}
    byp = collections.defaultdict(list)
    for k in ids:
        p, n = k.rsplit("-", 1)
        byp[p].append(int(n))
    gaps = {p: [i for i in range(1, max(ns) + 1) if i not in ns] for p, ns in byp.items()}
    gaps = {p: g for p, g in gaps.items() if g}
    return dict(duplicates=dup, id_gaps=gaps,
                counts={sh["title"]: sum(1 for u in sh["ucs"] if u["id"]) for sh in sheets})


def render(path, sheets, src, hyg, ann):
    out = [f"# Kiểm kê UC — lượt 1 (trích xuất nguyên văn)\n",
           f"Nguồn: `{path.split('/')[-1]}` · Sinh bằng `scripts/extract_uc_list.py` · "
           "Trường **đậm** = nguyên văn nguồn; trường *nghiêng* = suy ra ở bước phân tích (xem analysis.md).\n",
           "## Số UC theo sheet\n", "| Sheet | Số UC |", "|---|---|"]
    out += [f"| {k} | {v} |" for k, v in hyg["counts"].items()]
    out.append(f"| **Tổng** | **{sum(hyg['counts'].values())}** |\n")
    for sh in sheets:
        out.append(f"## Sheet: {sh['title']}\n")
        if sh["others"]:
            out.append("Dòng có dữ liệu nhưng không phải UC hay tiêu đề section (chép nguyên văn):\n")
            out.append("| Dòng | Nội dung |")
            out.append("|---|---|")
            for r, cells in sh["others"]:
                txt = " · ".join(f"{h or '(không tiêu đề)'}: {v}" for h, v in cells).replace("|", "\\|")
                out.append(f"| R{r} | {txt} |")
            out.append("")
        if not sh["ucs"]:
            out.append("_Không có UC._\n")
            continue
        cur = object()
        for uc in sh["ucs"]:
            if uc["section"] != cur:
                cur = uc["section"]
                out.append(f"### {cur or '(không thuộc section)'}\n")
                sec = next((x for x in sh["sections"] if x[1] == cur), None)
                if sec and sec[2]:
                    out.append(f"`{sh['title']}!R{sec[0]}` · " + " · ".join(f"**{h}:** `{v}`" for h, v in sec[2]) + "\n")
            fields = dict(uc["fields"])
            name_key = next((h for h in sh["headers"] if h and "name" in h.lower()), None)
            title = uc["id"]
            out.append(f"#### {title} · {fields.get(name_key) or '—'}")
            out.append(f"`{sh['title']}!R{uc['row']}`\n")
            for h, v in uc["fields"]:
                if h is None or h == sh["headers"][0] or h == name_key:
                    continue
                out.append(f"- **{h}:** {v if v else '—'}")
            for flag in uc["flags"]:
                out.append(f"- *Cờ dữ liệu:* {flag}")
            for k, v in (ann.get(uc["id"] or "", {}) or {}).items():
                out.append(f"- *{k}:* {v}")
            out.append("")
    return "\n".join(out)


def coverage(wb, text):
    """Chiều ngược: mọi ô có dữ liệu trong nguồn phải có mặt nguyên văn trong inventory."""
    total, miss = 0, []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                v = norm(c.value)
                if not v:
                    continue
                total += 1
                if v not in text and v.replace("|", "\\|") not in text:
                    miss.append(f"{ws.title}!{c.coordinate}: {v[:60]}")
    return dict(cells=total, missing=miss)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx"); ap.add_argument("out")
    ap.add_argument("--annotations"); ap.add_argument("--summary")
    a = ap.parse_args()
    wb, sheets = extract(a.xlsx)
    rare_actors(sheets)
    src = inventory_source(wb)
    hyg = hygiene(sheets)
    ann = json.load(open(a.annotations, encoding="utf-8")) if a.annotations else {}
    text = render(a.xlsx, sheets, src, hyg, ann)
    cov = coverage(wb, text)
    line = (f"Độ phủ chiều ngược: {cov['cells'] - len(cov['missing'])}/{cov['cells']} ô có dữ liệu "
            "trong file gốc xuất hiện nguyên văn trong file này.\n")
    text = text.replace("## Số UC theo sheet", line + "\n## Số UC theo sheet", 1)
    open(a.out, "w", encoding="utf-8").write(text)
    summary = dict(source=src, hygiene=hyg, coverage=cov,
                   flags=[dict(id=u["id"], where=f'{u["sheet"]}!R{u["row"]}', flags=u["flags"])
                          for sh in sheets for u in sh["ucs"] if u["flags"]])
    if a.summary:
        json.dump(summary, open(a.summary, "w", encoding="utf-8"), ensure_ascii=False, indent=2)
    print(json.dumps(summary, ensure_ascii=False, indent=1))


if __name__ == "__main__":
    sys.exit(main())
