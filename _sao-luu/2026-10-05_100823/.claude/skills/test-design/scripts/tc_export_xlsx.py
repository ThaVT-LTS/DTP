#!/usr/bin/env python3
"""
Xuất test case (markdown theo mẫu references/templates.md §2) sang Excel để chạy tay hoặc nhập công cụ quản lý test.

Dùng:
  python tc_export_xlsx.py tests/cases/*.md --out tests/test-cases.xlsx [--uc-pattern "UC-[A-Z]+-\\d{3}"]

Sheet:
  Test cases  — mỗi TC một dòng, nhóm theo UC → I. Redirect / II. Validation / III. Function;
                cột P–U (nền vàng) để người chạy test điền.
  Tổng hợp    — số TC và kết quả chạy theo UC, tính bằng công thức COUNTIFS (tự cập nhật khi điền trạng thái).
  Hướng dẫn   — cột nào cần điền, giá trị hợp lệ, một dòng ví dụ.
"""
import argparse, math, re
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

UC = r"UC-[A-Z]+-\d{3}"
GROUP_ORDER = [("R", "Redirect", "I. REDIRECT"), ("V", "Validation", "II. VALIDATION"), ("F", "Function", "III. FUNCTION")]
STATUSES = ["Pass", "Fail", "Blocked", "Not Run", "N/A"]
COLS = [  # (tiêu đề, độ rộng, khóa trường trong md)
    ("STT", 5, None), ("UC", 13, None), ("TC ID", 20, None), ("Nhóm", 11, "Nhóm"), ("Tiêu đề", 36, None),
    ("Loại", 10, "Loại"), ("Ưu tiên", 8, "Ưu tiên"), ("Kỹ thuật", 14, "Kỹ thuật"), ("Nguồn", 26, "Nguồn"),
    ("Tiền điều kiện", 32, "Tiền điều kiện"), ("Dữ liệu", 30, "Dữ liệu"), ("Bước", 38, "Bước"),
    ("Kết quả mong đợi", 46, "Mong đợi"), ("Tag", 22, "Tag"), ("Tự động hóa", 14, "Tự động hóa"),
    ("Kết quả thực tế", 30, None), ("Trạng thái chạy", 12, None), ("Ngày chạy", 12, None),
    ("Người chạy", 14, None), ("Bug ID", 12, None), ("Ghi chú", 24, None),
]
FIRST_INPUT_COL = 16  # P

FONT = Font(name="Arial", size=10)
BOLD = Font(name="Arial", size=10, bold=True)
WHITE_BOLD = Font(name="Arial", size=10, bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="1F4E78")
UC_FILL = PatternFill("solid", fgColor="BDD7EE")
GROUP_FILL = PatternFill("solid", fgColor="DDEBF7")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")


def parse(files, uc_re):
    tc_head = re.compile(rf"^#{{3,4}}\s+({uc_re}-TC-([RVF])\d+)\s*[·:\-–]\s*(.+?)\s*$")
    uc_head = re.compile(rf"^##\s+({uc_re})\s*[·:\-–]?\s*(.*?)\s*$")
    ucs, tcs, cur = {}, [], None
    for f in files:
        for line in open(f, encoding="utf-8"):
            line = line.rstrip("\n")
            m = tc_head.match(line)
            if m:
                cur = dict(id=m.group(1), letter=m.group(2), title=m.group(3), uc=m.group(1).split("-TC-")[0], f={})
                tcs.append(cur)
                continue
            m = uc_head.match(line)
            if m:
                ucs[m.group(1)] = m.group(2)
            if re.match(r"^#{1,6}\s", line):
                cur = None
                continue
            if cur is not None:
                fm = re.match(r"^- ([^:]+?):\s*(.*)$", line)
                if fm:
                    cur["f"][fm.group(1).strip()] = fm.group(2).strip()
    return ucs, tcs


def pretty(key, text):
    if not text:
        return ""
    if key == "Bước":
        return re.sub(r";\s*(?=\d+\))", "\n", text)
    if key in ("Mong đợi", "Tiền điều kiện", "Dữ liệu"):
        parts = [p.strip() for p in text.split(";") if p.strip()]
        return "\n".join(("- " + p) if key == "Mong đợi" and len(parts) > 1 else p for p in parts)
    return text


def row_height(values):
    lines = 1
    for (_, width, _), v in zip(COLS, values):
        if not v:
            continue
        chars = max(1, int(width * 1.1))
        n = sum(max(1, math.ceil(len(seg) / chars)) for seg in str(v).split("\n"))
        lines = max(lines, n)
    return min(409, 13.5 * lines + 3)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("cases", nargs="+")
    ap.add_argument("--out", required=True)
    ap.add_argument("--uc-pattern", default=UC)
    a = ap.parse_args()
    ucs, tcs = parse(a.cases, a.uc_pattern)
    tcs = [t for t in tcs if "deprecated" not in t["f"].get("Trạng thái", "").lower()]

    wb = Workbook()
    ws = wb.active
    ws.title = "Test cases"
    for c, (name, width, _) in enumerate(COLS, 1):
        cell = ws.cell(row=1, column=c, value=name)
        cell.font, cell.fill, cell.border = WHITE_BOLD, HEAD_FILL, BORDER
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        ws.column_dimensions[get_column_letter(c)].width = width
    ws.row_dimensions[1].height = 30

    dv = DataValidation(type="list", formula1='"' + ",".join(STATUSES) + '"', allow_blank=True,
                        error="Chọn một giá trị: " + ", ".join(STATUSES), errorTitle="Trạng thái không hợp lệ")
    ws.add_data_validation(dv)

    def section(r, text, fill, bold=True):
        for c in range(1, len(COLS) + 1):
            cell = ws.cell(row=r, column=c)
            cell.fill, cell.border = fill, BORDER
        cell = ws.cell(row=r, column=3, value=text)
        cell.font = BOLD if bold else FONT
        ws.row_dimensions[r].height = 18

    r, stt = 2, 0
    uc_order = list(dict.fromkeys([t["uc"] for t in tcs]))
    for u in uc_order:
        section(r, f"{u} · {ucs.get(u, '')}".strip(" ·"), UC_FILL)
        r += 1
        for letter, _, label in GROUP_ORDER:
            group = [t for t in tcs if t["uc"] == u and t["letter"] == letter]
            if not group:
                continue
            section(r, label, GROUP_FILL)
            r += 1
            for t in group:
                stt += 1
                vals = []
                for name, _, key in COLS:
                    if name == "STT":
                        vals.append(stt)
                    elif name == "UC":
                        vals.append(t["uc"])
                    elif name == "TC ID":
                        vals.append(t["id"])
                    elif name == "Tiêu đề":
                        vals.append(t["title"])
                    elif name == "Trạng thái chạy":
                        vals.append("Not Run")
                    elif key:
                        vals.append(pretty(key, t["f"].get(key, "")))
                    else:
                        vals.append(None)
                for c, v in enumerate(vals, 1):
                    cell = ws.cell(row=r, column=c, value=v)
                    cell.font, cell.border, cell.alignment = FONT, BORDER, WRAP
                    if c >= FIRST_INPUT_COL:
                        cell.fill = INPUT_FILL
                ws.row_dimensions[r].height = row_height(vals)
                r += 1
    last = r - 1
    status_col = get_column_letter([n for n, _, _ in COLS].index("Trạng thái chạy") + 1)
    dv.add(f"{status_col}2:{status_col}{max(last, 2)}")
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLS))}{max(last, 1)}"

    # ---------- Tổng hợp ----------
    ss = wb.create_sheet("Tổng hợp")
    ss["A1"] = "Tổng hợp test case và kết quả chạy"
    ss["A1"].font = Font(name="Arial", size=12, bold=True)
    ss["A2"] = "Mọi số liệu tính bằng công thức từ sheet 'Test cases' — tự cập nhật khi điền cột Trạng thái chạy."
    ss["A2"].font = Font(name="Arial", size=9, italic=True)
    heads = ["UC", "Tổng TC", "Redirect", "Validation", "Function", "P1", "Pass", "Fail", "Blocked", "Not Run", "N/A",
             "% đã chạy", "% Pass"]
    for c, h in enumerate(heads, 1):
        cell = ss.cell(row=4, column=c, value=h)
        cell.font, cell.fill, cell.border = WHITE_BOLD, HEAD_FILL, BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ss.column_dimensions[get_column_letter(c)].width = 14 if c > 1 else 18
    TCS = "'Test cases'!"
    rng = lambda col: f"{TCS}${col}$2:${col}${max(last, 2)}"
    UCR, GRR, PRR, STR = rng("B"), rng("D"), rng("G"), rng(status_col)
    row = 5
    for u in uc_order:
        A = f"$A{row}"
        f = {
            "Tổng TC": f"=COUNTIFS({UCR},{A})",
            "Redirect": f'=COUNTIFS({UCR},{A},{GRR},"Redirect")',
            "Validation": f'=COUNTIFS({UCR},{A},{GRR},"Validation")',
            "Function": f'=COUNTIFS({UCR},{A},{GRR},"Function")',
            "P1": f'=COUNTIFS({UCR},{A},{PRR},"P1")',
        }
        for s in STATUSES:
            f[s] = f'=COUNTIFS({UCR},{A},{STR},"{s}")'
        f["% đã chạy"] = f"=IFERROR((G{row}+H{row})/(B{row}-K{row}),0)"
        f["% Pass"] = f"=IFERROR(G{row}/(G{row}+H{row}),0)"
        ss.cell(row=row, column=1, value=u)
        for c, h in enumerate(heads[1:], 2):
            ss.cell(row=row, column=c, value=f[h])
        row += 1
    tot = row
    ss.cell(row=tot, column=1, value="Tổng")
    for c in range(2, 12):
        L = get_column_letter(c)
        ss.cell(row=tot, column=c, value=f"=SUM({L}5:{L}{tot - 1})")
    ss.cell(row=tot, column=12, value=f"=IFERROR((G{tot}+H{tot})/(B{tot}-K{tot}),0)")
    ss.cell(row=tot, column=13, value=f"=IFERROR(G{tot}/(G{tot}+H{tot}),0)")
    for rr in range(5, tot + 1):
        for c in range(1, len(heads) + 1):
            cell = ss.cell(row=rr, column=c)
            cell.font = BOLD if rr == tot else FONT
            cell.border = BORDER
            if c >= 12:
                cell.number_format = "0.0%"
    ss.cell(row=tot + 2, column=1, value="% đã chạy = (Pass + Fail) / (Tổng − N/A). % Pass = Pass / (Pass + Fail). Blocked và Not Run tính là chưa chạy.").font = Font(name="Arial", size=9, italic=True)

    # ---------- Hướng dẫn ----------
    hs = wb.create_sheet("Hướng dẫn")
    rows = [
        ("Hướng dẫn sử dụng", None),
        ("Nguồn", ", ".join(a.cases) + " — sinh bằng scripts/tc_export_xlsx.py. Sửa TC ở file markdown rồi xuất lại; chỉ điền kết quả chạy trong Excel."),
        ("Cột cần điền", "Cột P–U trên sheet 'Test cases' (nền vàng): Kết quả thực tế, Trạng thái chạy, Ngày chạy, Người chạy, Bug ID, Ghi chú."),
        ("Trạng thái chạy", "Pass · Fail · Blocked (không chạy được vì lỗi khác / thiếu môi trường) · Not Run (mặc định) · N/A (không áp dụng ở đợt này)."),
        ("Loại", "Positive = happy case · Negative = cố tình sai (dữ liệu sai, vi phạm rule, không quyền) · Abnormal = sự cố (mất mạng, bấm hai lần, đồng thời, hết phiên)."),
        ("Ưu tiên", "P1 luồng chính, phân quyền, dữ liệu, bảo mật · P2 luồng thay thế, rule, validation chính · P3 hiển thị, trường hợp hiếm."),
        ("Ví dụ một dòng đã chạy", "Kết quả thực tế: Hiện \"Họ và tên từ 2 đến 50 ký tự.\" · Trạng thái chạy: Pass · Ngày chạy: 2026-10-05 · Người chạy: Lan · Bug ID: (trống) · Ghi chú: (trống)"),
        ("Ví dụ dòng lỗi", "Kết quả thực tế: Tạo được người dùng với tên 51 ký tự · Trạng thái chạy: Fail · Ngày chạy: 2026-10-05 · Người chạy: Lan · Bug ID: BUG-123"),
    ]
    hs.column_dimensions["A"].width = 24
    hs.column_dimensions["B"].width = 110
    for i, (k, v) in enumerate(rows, 1):
        ka = hs.cell(row=i, column=1, value=k)
        ka.font = Font(name="Arial", size=12, bold=True) if v is None else BOLD
        ka.alignment = WRAP
        if v is not None:
            vb = hs.cell(row=i, column=2, value=v)
            vb.font, vb.alignment = FONT, WRAP
            hs.row_dimensions[i].height = 30

    wb.active = 0
    wb.save(a.out)
    print(f"{len(tcs)} TC → {a.out}")


if __name__ == "__main__":
    main()
