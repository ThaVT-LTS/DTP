#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Dựng file test case Excel từ bảng phẳng TSV.

    python tc_build_xlsx.py --data tc_login.tsv --sheet Login \
        --function "AUTH-L0001_Đăng nhập" --screen "Login · Customer Portal" \
        --creator "ThaVT" --source "Figma ... trang 1, 01/10/2026" --out TestCase.xlsx

TSV có header: I, II, III, IV, Pri, Name, Pre, Step, Expected, UC, Wait
Cột cha lặp lại ở mọi dòng — script tự gộp theo quy tắc trong references/format.md.
Thêm --append để ghi thêm một sheet vào file đã có.
"""
from __future__ import annotations
import argparse, csv, sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

FONT = "Times New Roman"
SIZE = 10

COLS = [
    ("No.", 6), ("I", 22), ("II", 24), ("III", 22), ("IV", 24),
    ("Priority", 9), ("TVP_ID", 11), ("Test case name/ Test view point", 44),
    ("Pre-condition", 26), ("Step", 40), ("Expected result", 52), ("Dev test", 10),
    ("Result", 11), ("Test person", 12), ("Test date", 12), ("Bug_ID", 9), ("Remark", 16),
    ("Result", 11), ("Test person", 12), ("Test date", 12), ("Bug_ID", 9), ("Remark", 16),
    ("Open question", 32), ("UC", 12),
]
GRP_ROW, HDR_ROW, DATA_ROW = 11, 12, 13
CLS = (2, 3, 4, 5)                       # B..E
PRI_OK = {"High", "Medium", "Low"}

# Bảng màu duy nhất — khớp references/format.md §4 và PALETTE trong tc_check.py
ink   = Font(name=FONT, size=SIZE)
bold  = Font(name=FONT, size=SIZE, bold=True)
white = Font(name=FONT, size=SIZE, bold=True, color="FFFFFF")
red   = Font(name=FONT, size=SIZE, color="B8362B")
head  = PatternFill("solid", fgColor="17324D")
subtle= PatternFill("solid", fgColor="EAE7E1")
pink  = PatternFill("solid", fgColor="FBE9E7")
thin  = Side(style="thin", color="C9C5BE")
BOX   = Border(left=thin, right=thin, top=thin, bottom=thin)
TOP   = Alignment(horizontal="left", vertical="top", wrap_text=True)
MID   = Alignment(horizontal="left", vertical="center", wrap_text=True)
CTR   = Alignment(horizontal="center", vertical="center", wrap_text=True)


def read_tsv(path: Path) -> list[dict]:
    with path.open(encoding="utf-8-sig", newline="") as f:
        rows = list(csv.DictReader(f, delimiter="\t"))
    need = ["I", "II", "III", "IV", "Pri", "Name", "Pre", "Step", "Expected", "UC"]
    missing = [c for c in need if rows and c not in rows[0]]
    if missing:
        sys.exit("Thiếu cột trong TSV: " + ", ".join(missing))
    for r in rows:
        for k, v in list(r.items()):
            r[k] = (v or "").replace("\\n", "\n").strip()
    return rows


def runs(ws, col, lo, hi):
    """Cụm dòng liền nhau cùng giá trị; ô trống = giống dòng trên."""
    vals, last = [], None
    for r in range(lo, hi + 1):
        v = ws.cell(r, col).value
        if v in (None, ""):
            v = last
        else:
            last = v
        vals.append(v)
    out, s = [], 0
    for i in range(1, len(vals) + 1):
        if i == len(vals) or vals[i] != vals[s]:
            if vals[s] is not None:
                out.append((lo + s, lo + i - 1))
            s = i
    return out


def build_sheet(wb, name, rows, meta):
    if name in wb.sheetnames:
        del wb[name]
    ws = wb.create_sheet(name)
    ws.sheet_view.showGridLines = False
    for i, (_, w) in enumerate(COLS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # khối thông tin
    info = [("Function Name", meta["function"]), ("Screen Name", meta["screen"]),
            ("Creator", meta["creator"]), ("Created date", meta["created"]),
            ("Reference", meta["source"])]
    for i, (k, v) in enumerate(info):
        a = ws.cell(3 + i, 2, k); a.font = bold; a.alignment = TOP; a.border = BOX
        b = ws.cell(3 + i, 3, v); b.font = ink;  b.alignment = TOP; b.border = BOX

    # bảng tổng kết theo vòng
    last = DATA_ROW + len(rows) - 1
    ws.cell(3, 11, "Round").font = bold
    for j, h in enumerate(["Total", "Passed", "Failed", "Pending", "Blocked", "NA", "Not execute"]):
        c = ws.cell(3, 12 + j, h); c.font = bold; c.alignment = CTR; c.border = BOX; c.fill = subtle
    for k, (lbl, col) in enumerate((("Round 1", "M"), ("Round 2", "R"))):
        r = 4 + k
        c = ws.cell(r, 11, lbl); c.font = bold; c.border = BOX
        ws.cell(r, 12, f"=COUNTA($G${DATA_ROW}:$G${last})")
        for j, st in enumerate(["Passed", "Failed", "Pending", "Blocked", "NA", "Not execute"]):
            ws.cell(r, 13 + j, f'=COUNTIF({col}{DATA_ROW}:{col}{last},"{st}")')
        for j in range(7):
            cc = ws.cell(r, 12 + j); cc.font = ink; cc.alignment = CTR; cc.border = BOX

    # tiêu đề hai tầng
    def grp(c0, c1, text):
        c = ws.cell(GRP_ROW, c0, text); c.font = white; c.fill = head; c.alignment = CTR
        ws.merge_cells(start_row=GRP_ROW, start_column=c0, end_row=GRP_ROW, end_column=c1)
        for i in range(c0, c1 + 1):
            ws.cell(GRP_ROW, i).border = BOX; ws.cell(GRP_ROW, i).fill = head
    grp(2, 5, "Classification"); grp(13, 17, "Test Round 1"); grp(18, 22, "Test Round 2")
    for i, (h, _) in enumerate(COLS, start=1):
        if i in (2, 3, 4, 5):
            c = ws.cell(HDR_ROW, i, h)
        elif i in range(13, 23):
            c = ws.cell(HDR_ROW, i, h)
        else:
            c = ws.cell(GRP_ROW, i, h)
            ws.merge_cells(start_row=GRP_ROW, start_column=i, end_row=HDR_ROW, end_column=i)
        c.font = white; c.fill = head; c.alignment = CTR; c.border = BOX
    for i in range(1, len(COLS) + 1):
        for r in (GRP_ROW, HDR_ROW):
            ws.cell(r, i).border = BOX
            if ws.cell(r, i).fill.fgColor.rgb in (None, "00000000"):
                ws.cell(r, i).fill = head
    ws.row_dimensions[GRP_ROW].height = 24
    ws.row_dimensions[HDR_ROW].height = 18

    # dữ liệu
    bad = []
    for i, row in enumerate(rows):
        r = DATA_ROW + i
        wait = row.get("Wait", "")
        if row["Pri"] and row["Pri"] not in PRI_OK:
            bad.append(f"dòng {r}: Priority '{row['Pri']}' không hợp lệ")
        vals = {1: i + 1, 2: row["I"], 3: row["II"], 4: row["III"], 5: row["IV"],
                6: row["Pri"], 7: "TC_%03d" % (i + 1), 8: row["Name"], 9: row["Pre"],
                10: row["Step"], 11: row["Expected"], 23: wait,
                24: ("UC-" + row["UC"]) if row["UC"] and not row["UC"].startswith("UC-") else row["UC"]}
        for c in range(1, len(COLS) + 1):
            cell = ws.cell(r, c, vals.get(c))
            cell.font = ink; cell.border = BOX; cell.alignment = TOP
            if c == 23 and wait:
                cell.font = red; cell.fill = pink
            if c == 11 and wait and not row["Expected"]:
                cell.fill = pink
        ws.row_dimensions[r].height = 32

    # gộp ô
    parent = [(DATA_ROW, last)]
    for col in CLS:
        rr = []
        for pa, pb in parent:
            rr += runs(ws, col, pa, pb)
        for a, b in rr:
            if b > a:
                ws.merge_cells(start_row=a, start_column=col, end_row=b, end_column=col)
                ws.cell(a, col).alignment = MID
        parent = rr or parent
    ws.freeze_panes = ws.cell(DATA_ROW, 8)
    return ws, bad


def main() -> int:
    ap = argparse.ArgumentParser(description="Dựng file test case Excel từ TSV phẳng.")
    ap.add_argument("--data", required=True)
    ap.add_argument("--sheet", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--function", default="")
    ap.add_argument("--screen", default="")
    ap.add_argument("--creator", default="")
    ap.add_argument("--created", default="")
    ap.add_argument("--source", default="")
    ap.add_argument("--append", action="store_true", help="ghi thêm sheet vào file đã có")
    a = ap.parse_args()

    rows = read_tsv(Path(a.data))
    if not rows:
        sys.exit("File dữ liệu rỗng.")
    out = Path(a.out)
    if a.append and out.is_file():
        wb = openpyxl.load_workbook(out)
    else:
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
    meta = dict(function=a.function or a.sheet, screen=a.screen, creator=a.creator,
                created=a.created, source=a.source)
    _, bad = build_sheet(wb, a.sheet, rows, meta)
    wb.save(out)
    print(f"{len(rows)} TC → {out} · sheet '{a.sheet}'")
    print(f"dòng còn câu hỏi mở: {sum(1 for r in rows if r.get('Wait'))}")
    for b in bad:
        print("CẢNH BÁO: " + b)
    print("Chạy tc_check.py trước khi bàn giao.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
