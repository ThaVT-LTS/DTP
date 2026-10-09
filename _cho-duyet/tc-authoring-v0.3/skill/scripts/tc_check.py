#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Kiểm tra file test case trước khi bàn giao.

    python tc_check.py TestCase.xlsx [--uc uc_list.txt] [--sheets Login,Scopes] [--like Login]

Báo lỗi (mã thoát 1):
  - TC thiếu UC_ID, hoặc TVP_ID trùng, hoặc thiếu Expected mà không khai ở cột Open question
  - Priority / Result không thuộc danh sách hợp lệ
  - Kết quả mong đợi mơ hồ ("hoạt động bình thường", "hiển thị đúng", "thành công"…)
  - Khoá {{...}} không có trong sheet _Data hoặc _Labels
  - Khu vực (Classification II) có thao tác ghi dữ liệu mà thiếu nhóm Không thành công
  - UC trong danh sách mà chưa có TC nào (khi truyền --uc)
  - Kiểu ô sai phông / cỡ chữ, chữ nghiêng hoặc gạch chân, màu nền / màu chữ ngoài bảng màu — ở mọi sheet,
    xét mọi kiểu ô đang dùng: ô có nội dung, ô trống đã định dạng (người chạy test sẽ gõ vào đó),
    kiểu mặc định của cột và dòng, chữ định dạng riêng từng đoạn trong một ô.
    Bảng màu lấy từ tc_palette.py; màu theme và màu indexed cũng bị xét, không được bỏ qua
  - Dòng có Open question mà ô đó chưa tô, hoặc đã hết câu hỏi mà vẫn tô
  - Khối tổng kết phía trên bảng đếm thiếu dòng dữ liệu (vùng đếm bắt đầu muộn hoặc kết thúc sớm)
  - Có --like: sheet khác form sheet mẫu trong cùng file (khối tiêu đề, độ rộng cột, trang in,
    kiểu ô dữ liệu, dropdown)
Cảnh báo: dòng còn chữ cái ngoài hệ Latin (chữ Nhật, Trung…) mà không khai ở cột Open question.
Dấu câu như · — – “ ” ← ⋮ không tính là chữ ngoài hệ Latin.
"""
from __future__ import annotations
import argparse, re, sys, unicodedata, warnings, zipfile
import xml.etree.ElementTree as ET
from pathlib import Path
import openpyxl
from openpyxl.styles.colors import COLOR_INDEX

# openpyxl báo nó không đọc dropdown và định dạng có điều kiện kiểu x14; script chỉ đọc, và phần đó đọc qua XML
warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")

from tc_palette import FONT, SIZE, PALETTE_FILL, PALETTE_FONT, OPEN_FILL

HDR_ROW, DATA_ROW = 11, 13               # hàng tiêu đề đầu tiên, dòng dữ liệu đầu tiên
C_NO, C_I, C_II, C_III, C_IV = 1, 2, 3, 4, 5
C_PRI, C_TVP, C_NAME, C_PRE, C_STEP, C_EXP = 6, 7, 8, 9, 10, 11
C_R1, C_R2, C_WAIT, C_UC = 13, 18, 23, 24
CLS_COLS = (C_I, C_II, C_III, C_IV)

PRI_OK = {"High", "Medium", "Low"}
RES_OK = {"Passed", "Failed", "Pending", "Blocked", "NA", "Not execute"}
VAGUE = ["hoạt động bình thường", "hiển thị đúng", "thành công", "bình thường",
         "đúng như mong đợi", "ok", "chính xác"]
KEY = re.compile(r"\{\{[^}]+\}\}")
SYS_SHEETS = {"Cover", "Hướng dẫn", "_Data", "_Labels", "_Patterns", "_Nav",
              "_Effects", "Test_Report", "Đối chiếu nhãn", "Cần xác nhận"}


def has_non_latin_letter(text: str) -> bool:
    """Chỉ bắt chữ cái của hệ chữ khác Latin; tiếng Việt có dấu vẫn là Latin."""
    return any(ch.isalpha() and not unicodedata.name(ch, "").startswith("LATIN") for ch in text)


def _color(color):
    """Mã RGB 6 ký tự của một màu (đối tượng openpyxl). Màu indexed được đổi ra RGB theo bảng chuẩn;
    màu theme trả 'theme:<n>'; chưa đặt màu thì trả None."""
    if color is None:
        return None
    kind = getattr(color, "type", None)
    if kind == "rgb":
        v = color.rgb
        return v[-6:].upper() if isinstance(v, str) else None
    if kind == "indexed":
        i = color.indexed
        if i == 64:
            return "000000"                    # màu chữ hệ thống
        if i == 65:
            return "FFFFFF"                    # màu nền cửa sổ
        return COLOR_INDEX[i][-6:].upper() if 0 <= i < len(COLOR_INDEX) else f"indexed:{i}"
    if kind == "theme":
        return f"theme:{color.theme}"
    return None


# ----------------------------------------------------------------------------- đọc thẳng XML
M = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
XM_SQREF = "{http://schemas.microsoft.com/office/excel/2006/main}sqref"
CELL = re.compile(r"([A-Z]+)(\d+)$")
RANGE = re.compile(r"\$?([A-Z]{1,3})\$?(\d+):\$?([A-Z]{1,3})\$?(\d+)")
KINDS = ("phông", "cỡ chữ", "kiểu chữ", "màu chữ", "màu nền")


def col_no(letters: str) -> int:
    n = 0
    for ch in letters:
        n = n * 26 + ord(ch) - 64
    return n


def col_name(n: int) -> str:
    s = ""
    while n:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


def _on(el) -> bool:
    """Một thuộc tính bật/tắt của phông (<i/>, <u/>…) có đang bật không."""
    return el is not None and el.get("val", "1") not in ("0", "false", "none")


class Book:
    """File .xlsx đọc thẳng từ XML. openpyxl bỏ qua dropdown x14, kiểu của ô trống và ô bị gộp,
    chữ định dạng riêng từng đoạn, và công thức (khi đọc giá trị) — nên các phép kiểm về form dùng lớp này."""

    def __init__(self, path: Path):
        z = zipfile.ZipFile(path)
        rels = [(r.get("Id"), r.get("Type", ""), r.get("Target"))
                for r in ET.fromstring(z.read("xl/_rels/workbook.xml.rels"))]

        def part(target):
            return target.lstrip("/") if target.startswith("/") else "xl/" + target

        by_id = {i: part(t) for i, _, t in rels}
        self.sheets = {}
        for sh in ET.fromstring(z.read("xl/workbook.xml")).iter(M + "sheet"):
            root = ET.fromstring(z.read(by_id[sh.get(R + "id")]))
            if root.find(M + "sheetData") is not None:                 # bỏ qua sheet biểu đồ
                self.sheets[sh.get("name")] = root
        st = ET.fromstring(z.read(next((part(t) for _, ty, t in rels if ty.endswith("/styles")), "xl/styles.xml")))
        self.fonts = list(st.find(M + "fonts"))
        self.fills = list(st.find(M + "fills"))
        self.xfs = list(st.find(M + "cellXfs"))
        custom = st.find(f"{M}colors/{M}indexedColors")               # file cũ có thể mang bảng màu indexed riêng
        self.indexed = ([c.get("rgb", "")[-6:].upper() for c in custom] if custom is not None
                        else [c[-6:].upper() for c in COLOR_INDEX])
        sst = next((part(t) for _, ty, t in rels if ty.endswith("/sharedStrings")), None)
        self.strings = ET.fromstring(z.read(sst)).findall(M + "si") if sst in z.namelist() else []
        self._faults: dict[str, list] = {}

    def color(self, el):
        """Mã RGB 6 ký tự của một phần tử màu. Màu indexed đổi ra RGB; màu theme trả 'theme:<n>';
        chưa đặt màu hoặc màu tự động thì trả None."""
        if el is None:
            return None
        if el.get("rgb"):
            return el.get("rgb")[-6:].upper()
        if el.get("indexed") is not None:
            i = int(el.get("indexed"))
            if i == 64:
                return "000000"                # màu chữ hệ thống
            if i == 65:
                return "FFFFFF"                # màu nền cửa sổ
            return self.indexed[i] if 0 <= i < len(self.indexed) else f"indexed:{i}"
        if el.get("theme") is not None:
            return f"theme:{el.get('theme')}"
        return None

    def fill(self, s):
        """Màu nền nhìn thấy của kiểu ô `s`; None khi không nền hoặc nền trắng."""
        pf = self.fills[int(self.xfs[int(s)].get("fillId", "0"))].find(M + "patternFill")
        if pf is None:
            return "nền chuyển màu"
        if pf.get("patternType") in (None, "none"):
            return None
        fg = pf.find(M + "fgColor")
        c = self.color(fg)
        if fg is not None and fg.get("theme") == "0" and float(fg.get("tint") or 0) == 0:
            c = "FFFFFF"                       # màu sáng 1 của theme = trắng
        return None if c == "FFFFFF" else (c or pf.get("patternType"))

    def font_faults(self, font, name_tag="name"):
        """Các điểm sai của một phông — của kiểu ô (name_tag='name') hay của một đoạn chữ ('rFont')."""
        out = []
        name = font.find(M + name_tag)
        name = name.get("val") if name is not None else None
        if name != FONT:
            out.append(("phông", name or "mặc định"))
        sz = font.find(M + "sz")
        if sz is not None and float(sz.get("val")) != SIZE:
            out.append(("cỡ chữ", f"{float(sz.get('val')):g}"))
        if _on(font.find(M + "i")):
            out.append(("kiểu chữ", "nghiêng"))
        if _on(font.find(M + "u")):
            out.append(("kiểu chữ", "gạch chân"))
        c = self.color(font.find(M + "color"))
        if c not in PALETTE_FONT:
            out.append(("màu chữ", c or "mặc định"))
        return out

    def style_faults(self, s):
        """Các điểm sai của kiểu ô `s`: phông, cỡ, kiểu chữ, màu chữ, màu nền."""
        if s not in self._faults:
            out = self.font_faults(self.fonts[int(self.xfs[int(s)].get("fontId", "0"))])
            bg = self.fill(s)
            if bg is not None and bg not in PALETTE_FILL:
                out.append(("màu nền", bg))
            self._faults[s] = out
        return self._faults[s]

    def merges(self, name):
        return sorted(m.get("ref") for m in self.sheets[name].iter(M + "mergeCell"))

    def merged_away(self, name) -> set:
        """Các ô bị gộp không phải ô neo — không hiển thị nên không xét phông, màu."""
        out = set()
        for ref in self.merges(name):
            a, _, b = ref.partition(":")
            ma, mb = CELL.match(a), CELL.match(b)
            if not (ma and mb):
                continue
            c1, r1, c2, r2 = col_no(ma.group(1)), int(ma.group(2)), col_no(mb.group(1)), int(mb.group(2))
            out.update((r, c) for r in range(r1, r2 + 1) for c in range(c1, c2 + 1) if (r, c) != (r1, c1))
        return out


def check_styles(book: Book, name: str, errors: list) -> None:
    """Mọi kiểu ô đang dùng trong sheet phải đúng phông, đúng cỡ, không nghiêng, không gạch chân,
    và chỉ dùng màu trong bảng màu. Gộp lỗi theo sheet."""
    root = book.sheets[name]
    bad = {k: [] for k in KINDS}
    hidden = book.merged_away(name)

    def note(where, faults, what=""):
        for kind, val in faults:
            bad[kind].append(f"{where}={what}{val}")

    for c in root.iter(M + "c"):
        m = CELL.match(c.get("r"))
        if (int(m.group(2)), col_no(m.group(1))) in hidden:
            continue
        note(c.get("r"), book.style_faults(c.get("s", "0")))
        text = None
        if c.get("t") == "s":
            v = c.find(M + "v")
            if v is not None and v.text and int(v.text) < len(book.strings):
                text = book.strings[int(v.text)]
        elif c.get("t") == "inlineStr":
            text = c.find(M + "is")
        if text is not None:                                           # chữ định dạng riêng từng đoạn
            for run in text.findall(M + "r"):
                rpr = run.find(M + "rPr")
                if rpr is not None:
                    note(c.get("r"), book.font_faults(rpr, "rFont"), "đoạn chữ ")
    for col in root.iter(M + "col"):                                   # kiểu mặc định của cột
        if col.get("style"):
            a, b = col_name(int(col.get("min"))), col_name(min(int(col.get("max")), 16384))
            note("cột " + (a if a == b else f"{a}:{b}"), book.style_faults(col.get("style")))
    for row in root.iter(M + "row"):                                   # kiểu mặc định của dòng
        if row.get("s") and row.get("customFormat") in ("1", "true"):
            note("dòng " + row.get("r"), book.style_faults(row.get("s")))
    for kind in KINDS:
        cells = bad[kind]
        if cells:
            more = f" … (+{len(cells) - 5})" if len(cells) > 5 else ""
            errors.append(f"{name} · {len(cells)} ô sai {kind}: {', '.join(cells[:5])}{more}")


def check_summary(book: Book, name: str, errors: list) -> None:
    """Khối tổng kết phía trên bảng phải đếm hết dữ liệu: mọi vùng trên cột TVP_ID / Result bắt đầu
    không muộn hơn dòng dữ liệu đầu và kết thúc không sớm hơn dòng dữ liệu cuối."""
    cols = {col_name(C_TVP), col_name(C_R1), col_name(C_R2)}
    last, formulas = 0, []
    for c in book.sheets[name].iter(M + "c"):
        m = CELL.match(c.get("r"))
        row, col = int(m.group(2)), col_no(m.group(1))
        if row >= DATA_ROW and col == C_TVP and c.find(M + "v") is not None:
            last = max(last, row)
        f = c.find(M + "f")
        if row < HDR_ROW and f is not None and f.text:
            formulas.append((c.get("r"), f.text))
    bad = []
    for ref, text in formulas:
        for m in RANGE.finditer(text):
            c1, r1, c2, r2 = m.group(1), int(m.group(2)), m.group(3), int(m.group(4))
            if c1 == c2 and c1 in cols and r2 >= DATA_ROW and (r1 > DATA_ROW or r2 < last):
                bad.append(f"{ref} đếm {c1}{r1}:{c2}{r2}")
                break
    if bad and last:
        more = f" … (+{len(bad) - 3})" if len(bad) > 3 else ""
        errors.append(f"{name} · {len(bad)} công thức tổng kết không phủ hết dữ liệu "
                      f"(dòng {DATA_ROW}–{last}): {', '.join(bad[:3])}{more}")


def collect_keys(wb) -> set[str]:
    keys = set()
    for name in ("_Data", "_Labels"):
        if name in wb.sheetnames:
            for row in wb[name].iter_rows(values_only=True):
                for v in row:
                    if isinstance(v, str):
                        keys.update(KEY.findall(v))
    return keys


def check_sheet(ws, keys, errors, warns):
    seen, ucs = set(), set()
    cur_ii = None
    areas = {}
    unmarked, stale = [], []
    for r in range(DATA_ROW, ws.max_row + 1):
        tvp = ws.cell(r, C_TVP).value
        if not tvp:
            continue
        where = f"{ws.title}!{r}"
        if tvp in seen:
            errors.append(f"{where} · TVP_ID trùng: {tvp}")
        seen.add(tvp)

        if ws.cell(r, C_II).value:
            cur_ii = ws.cell(r, C_II).value
        areas.setdefault(cur_ii, set()).add(str(ws.cell(r, C_III).value or ""))

        uc = ws.cell(r, C_UC).value
        if not uc:
            errors.append(f"{where} · {tvp} thiếu UC")
        else:
            ucs.add(str(uc).strip())

        pri = ws.cell(r, C_PRI).value
        if pri and pri not in PRI_OK:
            errors.append(f"{where} · Priority không hợp lệ: {pri}")
        for col, lbl in ((C_R1, "Round 1"), (C_R2, "Round 2")):
            v = ws.cell(r, col).value
            if v and v not in RES_OK:
                errors.append(f"{where} · {lbl} Result không hợp lệ: {v}")

        exp = (ws.cell(r, C_EXP).value or "").strip()
        wait = (ws.cell(r, C_WAIT).value or "").strip()
        if not exp and not wait:
            errors.append(f"{where} · {tvp} không có Expected và cũng không khai Open question")
        wc = ws.cell(r, C_WAIT)
        tinted = bool(wc.fill is not None and wc.fill.fill_type and _color(wc.fill.fgColor) == OPEN_FILL)
        if wait and not tinted:
            unmarked.append(tvp)
        elif tinted and not wait:
            stale.append(tvp)
        if exp:
            low = exp.lower()
            for w in VAGUE:
                if re.search(r"(^|[\s\-•])" + re.escape(w) + r"\s*$", low):
                    errors.append(f"{where} · {tvp} kết quả mong đợi mơ hồ: '{w}'")
                    break
        for cell in (C_PRE, C_STEP, C_EXP):
            txt = ws.cell(r, cell).value
            if isinstance(txt, str):
                for k in KEY.findall(txt):
                    if keys and k not in keys:
                        errors.append(f"{where} · khoá {k} chưa khai trong _Data/_Labels")
        if not wait:
            joined = " ".join(str(ws.cell(r, c).value or "") for c in (C_NAME, C_PRE, C_STEP, C_EXP))
            if has_non_latin_letter(joined):
                warns.append(f"{where} · {tvp} còn ký tự ngoài bảng chữ Latin mà không khai Open question")

    for label, ids in (("có Open question nhưng ô đó chưa tô", unmarked),
                       ("đã hết câu hỏi mà ô Open question vẫn tô", stale)):
        if ids:
            more = f" … (+{len(ids) - 5})" if len(ids) > 5 else ""
            errors.append(f"{ws.title} · {len(ids)} dòng {label}: {', '.join(ids[:5])}{more}")

    for area, lv3 in areas.items():
        if not area:
            continue
        joined = " ".join(lv3).lower()
        if "thành công" in joined and "không thành công" not in joined:
            errors.append(f"{ws.title} · khu vực '{area}' có nhóm Thành công nhưng thiếu Không thành công")
    return ucs


# ----------------------------------------------------------------------------- so với sheet mẫu
def compare_like(book: Book, like: str, targets: list[str]) -> list[str]:
    """Sheet TC phải cùng form với sheet mẫu `like` trong cùng file: khối tiêu đề cùng kiểu ô,
    cùng độ rộng cột, cùng cài đặt trang in, ô dữ liệu chỉ dùng kiểu sheet mẫu dùng cho cột đó,
    dropdown phủ hết dữ liệu."""
    if like not in book.sheets:
        return [f"Không có sheet mẫu '{like}'"]

    def flat(el):
        return (el.tag.split("}")[-1], tuple(sorted(el.attrib.items())), tuple(flat(c) for c in el))

    fonts = [flat(f) for f in book.fonts]
    fills = [flat(f) for f in book.fills]
    xfs = book.xfs

    def sig(s, loose=False, merged_col=False):
        """Chữ ký của một kiểu ô.
        loose=True bỏ nền và màu chữ — dòng còn câu hỏi mở được tô khác dòng thường.
        merged_col=True (bốn cột Classification) chỉ giữ phông, định dạng số, căn ngang, xuống dòng —
        viền và căn dọc của các cột này đổi theo vị trí ô trong cụm gộp."""
        xf = xfs[int(s)]
        font = fonts[int(xf.get("fontId", "0"))]
        if loose:
            font = (font[0], font[1], tuple(c for c in font[2] if c[0] != "color"))
        al = xf.find(M + "alignment")
        al = dict(al.attrib) if al is not None else {}
        if merged_col:
            return (font, xf.get("numFmtId"), al.get("horizontal"), al.get("wrapText"))
        return (font, None if loose else fills[int(xf.get("fillId", "0"))], xf.get("borderId"),
                xf.get("numFmtId"), tuple(sorted(al.items())))

    def load(name):
        root = book.sheets[name]
        cells = {}
        for c in root.iter(M + "c"):
            m = CELL.match(c.get("r"))
            cells[(int(m.group(2)), col_no(m.group(1)))] = c.get("s", "0")
        for rc in book.merged_away(name):            # ô bị gộp (không phải ô neo) không hiển thị → không so
            cells.pop(rc, None)
        cols = [(c.get("min"), c.get("max"), c.get("width"), sig(c.get("style", "0"), True))
                for c in root.iter(M + "col")]
        page = [(e.tag.split("}")[-1], tuple(sorted(e.attrib.items())))
                for tag in ("printOptions", "pageMargins", "pageSetup") for e in root.iter(M + tag)]
        dv = [(e.findtext(XM_SQREF) or e.get("sqref") or "")
              for e in root.iter() if e.tag.endswith("}dataValidation")]
        return cells, cols, book.merges(name), page, dv

    def head_merges(merges):
        return [m for m in merges if int(re.search(r"\d+", m).group()) < DATA_ROW]

    t_cells, t_cols, t_merges, t_page, t_dv = load(like)
    last_t = max((r for (r, c) in t_cells if r >= DATA_ROW and c == C_TVP), default=DATA_ROW)
    allowed = {}
    for (r, c), s in t_cells.items():
        if DATA_ROW <= r <= last_t:
            allowed.setdefault(c, set()).add(sig(s, True, c in CLS_COLS))
    errs = []
    for name in targets:
        if name == like or name not in book.sheets:
            continue
        cells, cols, merges, page, dv = load(name)
        bad = [f"{r}:{c}" for (r, c), s in t_cells.items()
               if r < DATA_ROW and (r, c) in cells and sig(cells[(r, c)]) != sig(s)]
        bad += [f"{r}:{c}" for (r, c) in t_cells if r < DATA_ROW and (r, c) not in cells]
        if bad:
            errs.append(f"{name} · khối tiêu đề khác sheet mẫu ở {len(bad)} ô, ví dụ hàng:cột {', '.join(bad[:5])}")
        if cols != t_cols:
            errs.append(f"{name} · độ rộng cột khác sheet mẫu '{like}'")
        if head_merges(merges) != head_merges(t_merges):
            errs.append(f"{name} · gộp ô ở khối tiêu đề khác sheet mẫu")
        if page != t_page:
            errs.append(f"{name} · cài đặt trang in khác sheet mẫu")
        last = max((r for (r, c) in cells if r >= DATA_ROW and c == C_TVP), default=DATA_ROW)
        odd = sorted({c for (r, c), s in cells.items()
                      if DATA_ROW <= r <= last and c in allowed
                      and sig(s, True, c in CLS_COLS) not in allowed[c]})
        if odd:
            errs.append(f"{name} · ô dữ liệu dùng kiểu sheet mẫu không có, ở cột: {', '.join(map(str, odd[:8]))}")
        if len(dv) != len(t_dv):
            errs.append(f"{name} · có {len(dv)} dropdown, sheet mẫu có {len(t_dv)}")
        for sq in dv:
            for tok in sq.split():
                m = re.match(r"[A-Z]+(\d+):[A-Z]+(\d+)$", tok)
                if m and int(m.group(1)) == DATA_ROW and int(m.group(2)) < last:
                    errs.append(f"{name} · dropdown {tok} không phủ hết dữ liệu (đến dòng {last})")
    return errs


def utf8_console() -> None:
    """Windows: khi kết quả in ra bị chuyển hướng (pipe, file), Python dùng bảng mã của máy và vấp chữ có dấu."""
    for stream in (sys.stdout, sys.stderr):
        if hasattr(stream, "reconfigure"):
            stream.reconfigure(encoding="utf-8", errors="replace")


def main() -> int:
    utf8_console()
    ap = argparse.ArgumentParser(description="Kiểm tra file test case.")
    ap.add_argument("xlsx")
    ap.add_argument("--uc", help="file danh sách UC, mỗi dòng một mã")
    ap.add_argument("--sheets", help="chỉ kiểm các sheet này, phân tách bằng dấu phẩy")
    ap.add_argument("--like", help="sheet TC mẫu trong cùng file; các sheet TC khác phải cùng form với nó")
    a = ap.parse_args()

    wb = openpyxl.load_workbook(Path(a.xlsx), data_only=True)
    book = Book(Path(a.xlsx))
    keys = collect_keys(wb)
    only = set(a.sheets.split(",")) if a.sheets else None

    errors, warns, ucs = [], [], set()
    n = 0
    for ws in wb.worksheets:
        if only and ws.title not in only:
            continue
        check_styles(book, ws.title, errors)
        if ws.title in SYS_SHEETS:
            continue
        n += 1
        ucs |= check_sheet(ws, keys, errors, warns)
        check_summary(book, ws.title, errors)
    if a.like:
        names = [ws.title for ws in wb.worksheets
                 if ws.title not in SYS_SHEETS and (not only or ws.title in only)]
        errors += compare_like(book, a.like, names)

    if a.uc:
        want = {l.strip() for l in Path(a.uc).read_text(encoding="utf-8").splitlines() if l.strip()}
        for u in sorted(want - ucs):
            errors.append(f"UC chưa có TC nào: {u}")

    print(f"{n} sheet · {len(ucs)} UC được phủ · {len(errors)} lỗi · {len(warns)} cảnh báo")
    for e in errors:
        print("LỖI: " + e)
    for w in warns:
        print("Cảnh báo: " + w)
    return 1 if errors else 0


if __name__ == "__main__":
    sys.exit(main())
